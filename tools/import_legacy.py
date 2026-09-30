#!/usr/bin/env python3
"""
Ubah ekspor Google Sheets (Tradewinds.xlsx) versi Apps Script menjadi satu file SQL impor.

  python3 tools/import_legacy.py Tradewinds.xlsx out.sql passwords.json

- out.sql       : tempel di Supabase > SQL Editor > Run. BERISI DATA PEMAIN - jangan commit ke repo.
- passwords.json: password sementara untuk pemain lama yang dulu login pakai Google (tanpa akun game).
                  Dibuat otomatis kalau belum ada, dan dipakai ulang di impor berikutnya.
Aman dijalankan ulang: kapten yang SUDAH diklaim pemiliknya tidak disentuh; yang belum, ditimpa data terbaru.
"""
import sys, json, re, uuid, hashlib, secrets, os, datetime
import openpyxl

NS = uuid.UUID('6f0c9a4e-3d1b-4c2a-9a5e-7b1d2c3e4f50')
WORDS = ['layar', 'jangkar', 'ombak', 'dermaga', 'kompas', 'karang', 'buritan', 'haluan', 'teluk', 'camar']


def legacy_hash(password, salt):  # sama dengan AuthService.gs hash_()
    h = salt + '|' + password
    for _ in range(250):
        h = hashlib.sha256((h + '|' + salt).encode('utf-8')).hexdigest()
    return h


def rows(wb, name):
    if name not in wb.sheetnames:
        return []
    it = list(wb[name].iter_rows(values_only=True))
    if not it:
        return []
    head = [str(h).strip() if h is not None else '' for h in it[0]]
    out = []
    for r in it[1:]:
        if all(c is None or c == '' for c in r):
            continue
        out.append({head[i]: r[i] for i in range(min(len(head), len(r))) if head[i]})
    return out


def q(v):
    if v is None:
        return 'null'
    return "'" + str(v).replace("'", "''") + "'"


def num(v, default=None):
    if v is None or v == '':
        return default
    try:
        f = float(v)
    except (TypeError, ValueError):
        return default
    return int(f) if f == int(f) else f


def sqlnum(v, default='null'):
    n = num(v, None)
    return default if n is None else str(n)


def ts(v):
    """Tanggal dari sheet (zona Asia/Jakarta) atau epoch ms -> literal timestamptz."""
    if v is None or v == '':
        return 'null'
    if isinstance(v, datetime.datetime):
        return q(v.strftime('%Y-%m-%d %H:%M:%S.%f') + '+07')
    n = num(v)
    if isinstance(n, (int, float)) and n > 1e11:
        return f"to_timestamp({n} / 1000.0)"
    return q(str(v) + '+07') if re.match(r'^\d{4}-\d\d-\d\d', str(v)) else 'null'


def js(v, default='{}'):
    if v is None or v == '':
        return q(default) + '::jsonb'
    try:
        json.loads(str(v))
        return q(str(v)) + '::jsonb'
    except Exception:
        return q(default) + '::jsonb'


def username_for(name, taken):
    u = re.sub(r'[^a-z0-9_]+', '_', str(name or '').lower()).strip('_')[:16] or 'kapten'
    if len(u) < 3:
        u = (u + '_kapten')[:16]
    base, i = u, 2
    while u in taken:
        u = (base[:14] + str(i)); i += 1
    taken.add(u)
    return u


def main(xlsx, out_sql, pw_file):
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    players = rows(wb, 'Players')
    accounts = {str(a['PlayerId']): a for a in rows(wb, 'Accounts') if a.get('PlayerId')}
    by = lambda sheet: {}
    group = {}
    for sheet in ['CharacterStats', 'Ship', 'PlayerLocation', 'PlayerInventory', 'PlayerBooks', 'PlayerMissions',
                  'PlayerLog', 'CombatLog', 'Warehouse', 'ShipEquipment']:
        for r in rows(wb, sheet):
            group.setdefault(sheet, {}).setdefault(str(r.get('PlayerId')), []).append(r)
    pw = json.load(open(pw_file)) if os.path.exists(pw_file) else {}

    L = ['-- Impor data Marantau versi Google Sheets -> Supabase',
         '-- Dibuat ' + datetime.datetime.now().strftime('%Y-%m-%d %H:%M') + ' dari ' + os.path.basename(xlsx),
         '-- BERISI DATA PEMAIN. Tempel di Supabase > SQL Editor > Run. Jangan commit ke GitHub.',
         'begin;', '']

    # 1) Konfigurasi dunia (termasuk jam dunia -> hari-game sama dengan versi lama)
    L.append('-- Konfigurasi')
    for c in rows(wb, 'GameConfig'):
        k, v = c.get('Key'), c.get('Value')
        if not k or v is None or v == '' or str(v) == 'None':
            continue
        if isinstance(v, bool):
            v = 'true' if v else 'false'
        elif isinstance(v, float) and v == int(v):
            v = int(v)
        L.append(f"select game.cfg_set({q(k)}, {q(v)});")

    # 2) Harga pasar
    L.append('\n-- Harga pasar')
    for m in rows(wb, 'Market'):
        b, cur = num(m.get('BasePrice')), num(m.get('CurrentPrice'))
        if b is None:
            continue
        L.append(f"update game.market set base_price = {round(b)}, current_price = {round(cur if cur is not None else b)} "
                 f"where city_id = {q(m['CityId'])} and commodity_id = {q(m['CommodityId'])};")

    # 3) Event kota yang sedang berjalan
    L.append('\n-- Event kota')
    L.append('delete from game.city_events;')
    for e in rows(wb, 'CityEvents'):
        L.append("insert into game.city_events(city_id, event_type, label, message, price_pct, reward_pct, expires_game_day) values "
                 f"({q(e['CityId'])}, {q(e['EventType'])}, {q(e.get('Label') or '')}, {q(e.get('Message') or '')}, "
                 f"{sqlnum(e.get('PriceMultiplierPercent'), '0')}, {sqlnum(e.get('RewardMultiplierPercent'), '0')}, "
                 f"{sqlnum(e.get('ExpiresGameDay'), '0')});")

    # 4) Kapten
    taken = {str(a['Username']).lower() for a in accounts.values()}
    summary = []
    for p in players:
        lid = str(p.get('PlayerId') or '')
        if not lid or not p.get('Archetype'):
            continue
        pid = str(uuid.uuid5(NS, lid))
        acc = accounts.get(lid)
        if acc:
            user, salt, phash, temp = str(acc['Username']).lower(), str(acc['Salt']), str(acc['PassHash']), None
        else:
            key = lid
            if key not in pw:
                pw[key] = {'username': username_for(p.get('CharacterName'), taken),
                           'password': secrets.choice(WORDS) + str(secrets.randbelow(9000) + 1000)}
            else:
                taken.add(pw[key]['username'])
            user, temp = pw[key]['username'], pw[key]['password']
            salt = str(uuid.uuid5(NS, 'salt:' + lid))
            phash = legacy_hash(temp, salt)
        summary.append((user, p.get('CharacterName'), temp))
        g = lambda s: group.get(s, {}).get(lid, [])
        B = []
        B.append(f"insert into game.players(player_id, username, character_name, archetype, gold, reputation, created_at, last_active, "
                 f"bank_balance, bank_last_interest_day, debt_balance, debt_last_interest_day, ship_upgrades, legacy_id) values "
                 f"('{pid}', {q(user)}, {q(p.get('CharacterName') or '')}, {q(p['Archetype'])}, {sqlnum(p.get('Gold'), '0')}, "
                 f"{js(p.get('Reputation'))}, coalesce({ts(p.get('CreatedAt'))}, now()), coalesce({ts(p.get('LastActive'))}, now()), "
                 f"{round(num(p.get('BankBalance'), 0))}, {sqlnum(p.get('BankLastInterestGameDay'))}, "
                 f"{round(num(p.get('DebtBalance'), 0))}, {sqlnum(p.get('DebtLastInterestGameDay'))}, {js(p.get('ShipUpgrades'))}, {q(lid)});")
        for s in g('CharacterStats')[:1]:
            B.append("insert into game.character_stats(player_id, trading, negotiation, navigation, sailing, combat, luck, knowledge) values "
                     f"('{pid}', " + ', '.join(sqlnum(s.get(k), d) for k, d in
                     [('Trading', '35'), ('Negotiation', '35'), ('Navigation', '35'), ('Sailing', '35'), ('Combat', '35'), ('Luck', '15'), ('Knowledge', '5')]) + ');')
        for s in g('Ship')[:1]:
            B.append("insert into game.ships(player_id, ship_name, tier, hull, max_hull, cargo, speed, combat, armor, navigation, condition, "
                     "condition_decay_per_sail, max_condition, damage_threshold) values "
                     f"('{pid}', {q(s.get('ShipName') or 'The Wandering Gull')}, {q(s.get('Tier') or 'I')}, "
                     + ', '.join(sqlnum(s.get(k), d) for k, d in
                     [('Hull', '100'), ('MaxHull', '100'), ('Cargo', '30'), ('Speed', '50'), ('Combat', '5'), ('Armor', '10'),
                      ('Navigation', '20'), ('Condition', '100'), ('ConditionDecayPerSail', '0.03'), ('MaxCondition', '100'),
                      ('DamageThreshold', '45')]) + ');')
        for s in g('PlayerLocation')[:1]:
            B.append("insert into game.player_location(player_id, city_id, arrived_game_day, destination_city_id, depart_at, arrive_at, pending_encounter) values "
                     f"('{pid}', {q(s.get('CurrentCityId') or 'sunda_empire')}, {sqlnum(s.get('ArrivedGameDay'))}, "
                     f"{q(s.get('DestinationCityId') or None)}, {ts(s.get('DepartAt'))}, {ts(s.get('ArriveAt'))}, "
                     f"{js(s.get('PendingEncounter'), 'null') if s.get('PendingEncounter') else 'null'});")
        inv = {}
        for s in g('PlayerInventory'):
            n = num(s.get('Qty'), 0) or 0
            if n > 0 and s.get('ItemId'):
                inv[str(s['ItemId'])] = inv.get(str(s['ItemId']), 0) + int(n)
        for item, n in inv.items():
            B.append(f"insert into game.inventory(player_id, item_id, qty) values ('{pid}', {q(item)}, {n});")
        for s in g('Warehouse'):
            n = num(s.get('Qty'), 0) or 0
            if n > 0:
                B.append(f"insert into game.warehouse(player_id, city_id, commodity_id, qty) values ('{pid}', {q(s['CityId'])}, {q(s['CommodityId'])}, {int(n)}) "
                         "on conflict do nothing;")
        for s in g('ShipEquipment'):
            if s.get('ItemId'):
                B.append(f"insert into game.ship_equipment(player_id, slot_type, item_id) values ('{pid}', {q(s.get('SlotType') or 'artifact')}, {q(s['ItemId'])}) on conflict do nothing;")
        for s in g('PlayerBooks'):
            if s.get('BookId'):
                B.append(f"insert into game.player_books(player_id, book_id, date_acquired) values ('{pid}', {q(s['BookId'])}, {sqlnum(s.get('DateAcquired'))}) on conflict do nothing;")
        active_seen = False
        for s in g('PlayerMissions'):
            st = str(s.get('Status') or 'completed')
            if st == 'active':
                if active_seen:
                    st = 'abandoned'
                active_seen = True
            B.append("insert into game.player_missions(player_id, mission_id, status, accepted_at, city_id, commodity_id, qty, deliver_to_city_id, "
                     "reward, type, source_city_id, loaded_qty) values "
                     f"('{pid}', {q(s.get('MissionId') or str(uuid.uuid4()))}, {q(st)}, {sqlnum(s.get('AcceptedAt'))}, {q(s.get('CityId'))}, "
                     f"{q(s.get('CommodityId') or '')}, {sqlnum(s.get('Qty'), '0')}, {q(s.get('DeliverToCityId'))}, {sqlnum(s.get('Reward'), '0')}, "
                     f"{q(s.get('Type') or 'legacy')}, {q(s.get('SourceCityId') or '')}, {sqlnum(s.get('LoadedQty'), '0')});")
        logs = g('PlayerLog')
        if logs:
            vals = [f"('{pid}', {sqlnum(s.get('GameDay'), '0')}, coalesce({ts(s.get('Timestamp'))}, now()), {q(s.get('Message') or '')})"
                    for s in logs if s.get('Message')]
            for i in range(0, len(vals), 200):
                B.append("insert into game.player_log(player_id, game_day, ts, message) values\n  " + ',\n  '.join(vals[i:i + 200]) + ';')
        cl = g('CombatLog')
        if cl:
            vals = [f"('{pid}', coalesce({ts(s.get('Timestamp'))}, now()), {sqlnum(s.get('EnemyLevel'))}, {q(s.get('Action'))}, {q(s.get('Result'))}, {q(s.get('Loot') or '')})"
                    for s in cl]
            B.append("insert into game.combat_log(player_id, ts, enemy_level, action, result, loot) values\n  " + ',\n  '.join(vals) + ';')
        B.append("insert into game.legacy_accounts(username, pass_hash, salt, player_id, legacy_id) values "
                 f"({q(user)}, {q(phash)}, {q(salt)}, '{pid}', {q(lid)}) on conflict (username) do update set "
                 "pass_hash = excluded.pass_hash, salt = excluded.salt, player_id = excluded.player_id, legacy_id = excluded.legacy_id, "
                 "failed = 0, imported_at = now();")

        cname = str(p.get('CharacterName') or '')
        L.append(f"\n-- Kapten: {cname} (username: {user})")
        L.append('do $imp$ begin')
        L.append(f"  if exists (select 1 from game.legacy_accounts where username = {q(user)} and claimed_at is not null) then")
        L.append(f"    raise notice 'Lewati %: sudah diklaim pemiliknya.', {q(user)}; return; end if;")
        L.append(f"  if exists (select 1 from game.players where username = {q(user)} and player_id <> '{pid}') then")
        L.append(f"    raise notice 'Lewati %: username sudah dipakai akun baru.', {q(user)}; return; end if;")
        L.append(f"  if exists (select 1 from game.players where lower(btrim(character_name)) = lower(btrim({q(cname)})) and player_id <> '{pid}') then")
        L.append(f"    raise notice 'Lewati %: nama kapten sudah dipakai akun baru.', {q(user)}; return; end if;")
        L.append(f"  delete from game.players where player_id = '{pid}';")
        L.extend('  ' + b for b in B)
        L.append('end $imp$;')

    L.append('\ncommit;')
    L.append("select username, legacy_id, claimed_at from game.legacy_accounts order by username;")
    open(out_sql, 'w').write('\n'.join(L) + '\n')
    json.dump(pw, open(pw_file, 'w'), indent=2)
    for u, n, t in summary:
        print(f"{u:18} {str(n):16} {'password sementara: ' + t if t else '(password lama)'}")


if __name__ == '__main__':
    if len(sys.argv) != 4:
        print(__doc__); sys.exit(1)
    main(*sys.argv[1:])
